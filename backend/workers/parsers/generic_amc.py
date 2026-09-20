"""Generic AMC monthly portfolio disclosure parser.

Conforms to BaseAMCParser contract Gates H1–H8.
Parses standard SEBI-mandated monthly portfolio workbooks (.xlsx).
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from datetime import date
from typing import Any

import openpyxl
from workers.holdings_worker import RawHoldingEntry
from workers.parsers.base import BaseAMCParser, ParsingResult


class GenericAMCParser(BaseAMCParser):
    """Generic parser for SEBI-compliant AMC monthly portfolio workbooks."""

    def __init__(self, amc_code: str = "GENERIC", default_scheme_filter: str | None = None):
        self.amc_code = amc_code.upper()
        self.default_scheme_filter = default_scheme_filter

    def parse_workbook(
        self,
        workbook_data: str | bytes | io.BytesIO,
        portfolio_id: int,
        as_of_date: date,
        disclosed_date: date,
        scheme_name_filter: str | None = None,
    ) -> ParsingResult:
        """Parse SEBI-mandated monthly Excel workbook into validated holdings list."""
        if isinstance(workbook_data, (bytes, bytearray)):
            wb = openpyxl.load_workbook(io.BytesIO(workbook_data), data_only=True)
        elif isinstance(workbook_data, io.BytesIO):
            wb = openpyxl.load_workbook(workbook_data, data_only=True)
        else:
            wb = openpyxl.load_workbook(workbook_data, data_only=True)

        target_filter = (scheme_name_filter or self.default_scheme_filter or "").strip().lower()
        target_sheet = None

        # 1. Match sheet by name if multiple sheets exist
        if target_filter:
            for name in wb.sheetnames:
                if target_filter in name.lower():
                    target_sheet = wb[name]
                    break

        if target_sheet is None:
            target_sheet = wb.active or wb.worksheets[0]

        rows_data: list[list[Any]] = []
        for row in target_sheet.iter_rows(values_only=True):
            rows_data.append(list(row))

        if not rows_data:
            return ParsingResult(
                holdings=[],
                total_weight=0.0,
                equity_count=0,
                valid=False,
                validation_message=f"Sheet '{target_sheet.title}' is empty",
            )

        # 2. Sniff header row (Gate H1)
        header_map = self.sniff_headers(rows_data)

        # 3. Parse data rows
        holdings: list[RawHoldingEntry] = []
        start_row = header_map.header_row_index + 1

        for row in rows_data[start_row:]:
            if not row or all(c is None for c in row):
                continue

            raw_name = row[header_map.name_col] if header_map.name_col < len(row) else None
            if not raw_name:
                continue

            name_str = str(raw_name).strip()
            upper_name = name_str.upper()

            # Skip subtotal, total, header, and footer rows
            if any(k in upper_name for k in ("TOTAL", "SUB TOTAL", "GRAND TOTAL", "PORTFOLIO AS ON", "NOTES:", "DISCLAIMER", "NET ASSETS")):
                continue

            raw_pct = row[header_map.pct_nav_col] if header_map.pct_nav_col < len(row) else None
            pct_nav = self.clean_weight(raw_pct)
            if pct_nav <= 0.0:
                continue

            raw_isin = row[header_map.isin_col] if header_map.isin_col < len(row) else None
            cleaned_isin = self.clean_isin(raw_isin)

            raw_sector = (
                row[header_map.sector_col]
                if header_map.sector_col is not None and header_map.sector_col < len(row)
                else None
            )
            sector = str(raw_sector).strip() if raw_sector else None

            raw_qty = (
                row[header_map.quantity_col]
                if header_map.quantity_col is not None and header_map.quantity_col < len(row)
                else None
            )
            try:
                quantity = float(raw_qty) if raw_qty is not None else None
            except (ValueError, TypeError):
                quantity = None

            raw_mval = (
                row[header_map.market_value_col]
                if header_map.market_value_col is not None and header_map.market_value_col < len(row)
                else None
            )
            try:
                mval = float(raw_mval) if raw_mval is not None else None
            except (ValueError, TypeError):
                mval = None

            asset_type = self.detect_asset_type(name_str, cleaned_isin, sector)

            holdings.append(
                RawHoldingEntry(
                    portfolio_id=portfolio_id,
                    as_of_date=as_of_date,
                    isin=cleaned_isin,
                    security_name=name_str,
                    asset_type=asset_type,
                    sector=sector,
                    quantity=quantity,
                    market_value_lakhs=mval,
                    pct_nav=pct_nav,
                    disclosed_date=disclosed_date,
                )
            )

        valid, msg = self.validate_holdings(holdings)
        total_w = round(sum(h.pct_nav for h in holdings), 4)
        eq_count = sum(1 for h in holdings if h.asset_type == "EQUITY")

        return ParsingResult(
            holdings=holdings,
            total_weight=total_w,
            equity_count=eq_count,
            valid=valid,
            validation_message=msg,
        )


class HDFCParser(GenericAMCParser):
    """Parser for HDFC Asset Management monthly portfolio disclosures."""

    def __init__(self, default_scheme_filter: str | None = "HDFC Small Cap"):
        super().__init__(amc_code="HDFC", default_scheme_filter=default_scheme_filter)


class ICICIPrudentialParser(GenericAMCParser):
    """Parser for ICICI Prudential AMC monthly portfolio disclosures."""

    def __init__(self, default_scheme_filter: str | None = "Bluechip"):
        super().__init__(amc_code="ICICI_PRU", default_scheme_filter=default_scheme_filter)


class SBIParser(GenericAMCParser):
    """Parser for SBI Funds Management monthly portfolio disclosures."""

    def __init__(self, default_scheme_filter: str | None = "Small Cap"):
        super().__init__(amc_code="SBI", default_scheme_filter=default_scheme_filter)


class KotakParser(GenericAMCParser):
    """Parser for Kotak Mahindra AMC monthly portfolio disclosures."""

    def __init__(self, default_scheme_filter: str | None = "Small Cap"):
        super().__init__(amc_code="KOTAK", default_scheme_filter=default_scheme_filter)


class QuantParser(GenericAMCParser):
    """Parser for Quant Mutual Fund monthly portfolio disclosures."""

    def __init__(self, default_scheme_filter: str | None = "Small Cap"):
        super().__init__(amc_code="QUANT", default_scheme_filter=default_scheme_filter)


class AxisParser(GenericAMCParser):
    """Parser for Axis Mutual Fund monthly portfolio disclosures."""

    def __init__(self, default_scheme_filter: str | None = "Small Cap"):
        super().__init__(amc_code="AXIS", default_scheme_filter=default_scheme_filter)


class PPFASParser(GenericAMCParser):
    """Parser for PPFAS Mutual Fund monthly portfolio disclosures."""

    def __init__(self, default_scheme_filter: str | None = "Flexi Cap"):
        super().__init__(amc_code="PPFAS", default_scheme_filter=default_scheme_filter)


class TataParser(GenericAMCParser):
    """Parser for Tata Mutual Fund monthly portfolio disclosures."""

    def __init__(self, default_scheme_filter: str | None = "Small Cap"):
        super().__init__(amc_code="TATA", default_scheme_filter=default_scheme_filter)


class BandhanParser(GenericAMCParser):
    """Parser for Bandhan Mutual Fund monthly portfolio disclosures."""

    def __init__(self, default_scheme_filter: str | None = "Small Cap"):
        super().__init__(amc_code="BANDHAN", default_scheme_filter=default_scheme_filter)


class InvescoParser(GenericAMCParser):
    """Parser for Invesco Mutual Fund monthly portfolio disclosures."""

    def __init__(self, default_scheme_filter: str | None = "Small Cap"):
        super().__init__(amc_code="INVESCO", default_scheme_filter=default_scheme_filter)


class DSPParser(GenericAMCParser):
    """Parser for DSP Mutual Fund monthly portfolio disclosures."""

    def __init__(self, default_scheme_filter: str | None = "Small Cap"):
        super().__init__(amc_code="DSP", default_scheme_filter=default_scheme_filter)

