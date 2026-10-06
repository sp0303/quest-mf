"""TER ingestion from AMFI's official "TER of MF schemes" feed (spec §15/§16.2 cost).

AMFI publishes the daily Total Expense Ratio of every scheme (Regular and Direct).
We store the Direct-plan TER (D_TER, in %) for each portfolio's canonical
Direct-Growth scheme in ref.ter_history, keeping only change points so the compute
worker can look it up point-in-time (Rule Q2). Unmatched schemes are skipped,
never guessed (genuine data only).

Usage:
    python -m workers.ter_worker            # current + previous month
    python -m workers.ter_worker --months 6
"""

from __future__ import annotations

import argparse
import asyncio
import io
import logging
import re
from datetime import date, datetime
from typing import Any

import asyncpg
import httpx
import openpyxl

from app.config import settings
from workers.ingestion_worker import AsyncRateLimiter, store_raw_payload

logger = logging.getLogger("ter_worker")
logging.basicConfig(level=logging.INFO)

AMFI_BASE = "https://www.amfiindia.com/api"
_rate = AsyncRateLimiter(requests_per_second=1.0)
_STOPWORDS = {"fund", "the", "scheme", "plan", "an", "open", "ended"}


def normalize_scheme_name(name: str) -> str:
    """Canonical key for matching AMFI TER names to ref.portfolios.display_name."""
    s = name.lower().replace("&", " and ")
    s = re.sub(r"\(.*?\)", " ", s)  # drop "(erstwhile ...)" / "(formerly ...)"
    s = re.sub(r"[^a-z0-9]+", " ", s)
    s = re.sub(r"\b(small|mid|large|flexi|multi)cap\b", r"\1 cap", s)  # "Smallcap" == "Small Cap"
    return " ".join(w for w in s.split() if w not in _STOPWORDS)


def financial_year(d: date) -> str:
    return f"{d.year}-{d.year + 1}" if d.month >= 4 else f"{d.year - 1}-{d.year}"


def _as_date(v: Any) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return datetime.fromisoformat(str(v).replace("Z", "+00:00")).date()


def parse_ter_workbook(data: bytes) -> list[dict[str, Any]]:
    """Rows of AMFI's TER Excel export as {Scheme_Name, TER_Date, D_TER} dicts."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    rows = wb[wb.sheetnames[0]].iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows, ())]
    try:
        i_name = header.index("scheme name")
        i_date = header.index("ter date")
        i_ter = header.index("direct plan - total ter (%)")
    except ValueError as exc:
        raise ValueError(f"Unexpected AMFI TER header: {header}") from exc
    return [
        {"Scheme_Name": r[i_name], "TER_Date": r[i_date], "D_TER": r[i_ter]}
        for r in rows
        if r and r[i_name]
    ]


def ter_change_points(rows: list[dict[str, Any]]) -> dict[str, list[tuple[date, float]]]:
    """Group AMFI rows by normalized name -> sorted [(date, D_TER%)] change points."""
    series: dict[str, dict[date, float]] = {}
    for r in rows:
        try:
            ter = float(r["D_TER"])
            d = _as_date(r["TER_Date"])
        except (KeyError, TypeError, ValueError):
            continue
        if ter <= 0:
            continue
        series.setdefault(normalize_scheme_name(r["Scheme_Name"]), {})[d] = ter
    out: dict[str, list[tuple[date, float]]] = {}
    for key, by_date in series.items():
        pts: list[tuple[date, float]] = []
        for d in sorted(by_date):
            if not pts or pts[-1][1] != by_date[d]:
                pts.append((d, by_date[d]))
        out[key] = pts
    return out


async def _get(client: httpx.AsyncClient, url: str) -> httpx.Response:
    await _rate.wait()
    resp = await client.get(url)
    resp.raise_for_status()
    return resp


async def fetch_month(client: httpx.AsyncClient, month_number: str) -> list[dict[str, Any]]:
    """One AMFI Excel export per month (all schemes, all days)."""
    url = (
        f"{AMFI_BASE}/populate-te-rdata-revised?MF_ID=All&Month={month_number}"
        "&strCat=-1&strType=-1&excel=true"
    )
    data = (await _get(client, url)).content
    store_raw_payload("amfi_ter", data, f"ter_{month_number}.xlsx")
    return parse_ter_workbook(data)


async def fetch_recent_months(n_months: int) -> list[dict[str, Any]]:
    headers = {"User-Agent": "Mozilla/5.0 (quest-mf TER ingestion)"}
    async with httpx.AsyncClient(timeout=120.0, headers=headers) as client:
        today = date.today()
        months: list[str] = []
        for fy in (financial_year(today), financial_year(date(today.year - 1, today.month, 1))):
            listing = (await _get(client, f"{AMFI_BASE}/populate-ter-month?year={fy}")).json()
            months += [m["MonthNumber"] for m in listing if m.get("MonthNumber") not in months]
        rows: list[dict[str, Any]] = []
        for m in months[:n_months]:
            month_rows = await fetch_month(client, m)
            logger.info("AMFI TER %s: %d rows", m, len(month_rows))
            rows.extend(month_rows)
        return rows


async def upsert_ter(conn: asyncpg.Connection, points: dict[str, list[tuple[date, float]]]) -> int:
    schemes = await conn.fetch(
        """
        SELECT s.scheme_code, p.display_name
        FROM ref.schemes s JOIN ref.portfolios p ON p.portfolio_id = s.portfolio_id
        WHERE s.is_canonical = true AND upper(s.plan) = 'DIRECT' AND upper(s.option) = 'GROWTH';
        """
    )
    records = [
        (r["scheme_code"], d, ter)
        for r in schemes
        for d, ter in points.get(normalize_scheme_name(r["display_name"]), [])
    ]
    matched = len({r[0] for r in records})
    await conn.executemany(
        """
        INSERT INTO ref.ter_history (scheme_code, effective_date, ter)
        VALUES ($1, $2, $3)
        ON CONFLICT (scheme_code, effective_date) DO UPDATE SET ter = EXCLUDED.ter;
        """,
        records,
    )
    logger.info(
        "TER: %d change points for %d/%d canonical schemes.", len(records), matched, len(schemes)
    )
    return matched


async def ingest_ter(n_months: int = 2) -> dict[str, int]:
    rows = await fetch_recent_months(n_months)
    points = ter_change_points(rows)
    conn = await asyncpg.connect(settings.pg_dsn)
    try:
        matched = await upsert_ter(conn, points)
    finally:
        await conn.close()
    return {"amfi_rows": len(rows), "amfi_schemes": len(points), "matched_schemes": matched}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Ingest AMFI TER (Direct plan) into ref.ter_history")
    ap.add_argument("--months", type=int, default=2)
    print(asyncio.run(ingest_ter(ap.parse_args().months)))
